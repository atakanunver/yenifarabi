import asyncio
import io
import sys
import types

import farabi_dosya_araci as fd
import pytest

BASLIK = ["Okul No", "T.C. Kimlik No", "Ad Soyad", "Sınıf", "Doğum Tarihi"]
SATIRLAR = [
    ["0123", "12345678901", "Ayşe Çelik", "9-A", "15.03.2010"],
    ["0124", "12345678902", "Mehmet Öz", "9-B", "02.01.2011"],
    ["0125", "12345678903", "Şule Ğür", "9-A", "20.12.2010"],
]


def test_tcsiz_tablo_turkce_ad_eslesmesiyle():
    b, s = fd.donustur(BASLIK, SATIRLAR, cikar="tc")
    assert b == ["Okul No", "Ad Soyad", "Sınıf", "Doğum Tarihi"]
    assert s[0] == ["0123", "Ayşe Çelik", "9-A", "15.03.2010"]  # başındaki 0 korunur


def test_sutun_secimi_sirayi_korur_ve_suzer():
    b, s = fd.donustur(BASLIK, SATIRLAR, sutunlar="ad soyad, okul no", filtre_sutun="sinif", filtre_deger="9-a")
    assert b == ["Ad Soyad", "Okul No"]
    assert s == [["Ayşe Çelik", "0123"], ["Şule Ğür", "0125"]]


def test_dogum_gunu_tablosu_ay_gune_gore():
    b, s = fd.donustur(BASLIK, SATIRLAR, sutunlar="ad soyad, doğum", dogum_gunune_gore=True)
    assert b == ["Ad Soyad", "Doğum Günü"]
    assert [r[1] for r in s] == ["02.01", "15.03", "20.12"]


def test_siralama_sayisal_ve_azalan():
    b, s = fd.donustur(["No"], [["9"], ["10"], ["2"]], sirala="no", azalan=True)
    assert [r[0] for r in s] == ["10", "9", "2"]


@pytest.mark.parametrize(
    "kw, parca",
    [({"cikar": "yok"}, "yok"), ({"sutunlar": "no"}, "birden fazla"), ({"cikar": "okul no, tc, ad, ad soyad, sınıf, doğum tarihi"}, "en az bir")],
)
def test_hatali_sutun_acik_mesaj(kw, parca):
    with pytest.raises(ValueError, match=parca):
        fd.donustur(["Ad", "Ad Soyad", "Okul No", "T.C. Kimlik No", "Sınıf", "Doğum Tarihi"], [], **kw)


def test_markdown_bloklari():
    bl = fd._markdown_bloklari("# Rapor\n\nMetin **kalın**.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n- madde\n")
    assert bl == [("h", 1, "Rapor"), ("p", "Metin kalın."), ("t", ["A", "B"], [["1", "2"]]), ("p", "• madde")]


def test_csv_noktalivirgul_ve_cp1254():
    veri = "Ad;Sınıf\nÇağla;9-A\n".encode("cp1254")
    assert fd._tablo_csv(veri) == (["Ad", "Sınıf"], [["Çağla", "9-A"]])


def test_xlsx_yaz_oku_gidis_donus():
    pytest.importorskip("openpyxl")
    veri = fd._xlsx([("Liste", BASLIK, SATIRLAR)])
    assert fd._tablo_xlsx(veri, "") == (BASLIK, SATIRLAR)


def test_docx_ve_pdf_uretilir():
    pytest.importorskip("docx")
    pytest.importorskip("fpdf")
    d = fd._docx("Başlık", [("t", BASLIK, SATIRLAR), ("p", "ğüşiöç")])
    assert fd._tablo_docx(d, "")[1][0][2] == "Ayşe Çelik"
    assert fd._pdf("Başlık", [("t", BASLIK, SATIRLAR), ("p", "ğüşiöç")])[:4] == b"%PDF"


class Dosya:
    def __init__(self, ad, yol):
        self.id, self.filename, self.path, self.user_id = "f1", ad, str(yol), "u1"


@pytest.fixture
def ortam(monkeypatch, tmp_path):
    kayit, kaydedilen = {}, []

    async def get_file_by_id(fid):
        return kayit.get(fid)

    mod = types.ModuleType("open_webui.models.files")
    mod.Files = types.SimpleNamespace(get_file_by_id=get_file_by_id)
    monkeypatch.setitem(sys.modules, "open_webui.models.files", mod)

    async def kaydet(veri, ad, mime, user):
        kaydedilen.append((ad, mime, veri))
        return f"id{len(kaydedilen)}"

    monkeypatch.setattr(fd.Tools, "_kaydet", staticmethod(kaydet))
    monkeypatch.setattr(fd.Tools, "_oku", staticmethod(lambda yol: open(yol, "rb").read()))
    return kayit, kaydedilen, tmp_path


def test_tablo_uret_uctan_uca(ortam):
    pytest.importorskip("openpyxl")
    kayit, kaydedilen, tmp = ortam
    yol = tmp / "ogrenciler.csv"
    yol.write_text(";".join(BASLIK) + "\n" + "\n".join(";".join(r) for r in SATIRLAR), encoding="utf-8")
    kayit["f1"] = Dosya("ogrenciler.csv", yol)
    t = fd.Tools()
    arg = {"__files__": [{"id": "f1", "type": "file"}], "__user__": {"id": "u1", "role": "user"}}
    bilgi = asyncio.run(t.tablo_bilgisi(**arg))
    assert "3 satır, 5 sütun" in bilgi and "T.C. Kimlik No" in bilgi
    cevap = asyncio.run(t.tablo_uret(cikar="tc", bicim="xlsx", baslik="TC siz liste", **arg))
    assert "/api/v1/files/id1/content?attachment=true" in cevap and "AYNEN" in cevap
    assert kaydedilen[0][0] == "TC_siz_liste.xlsx"
    assert fd._tablo_xlsx(kaydedilen[0][2], "")[0] == ["Okul No", "Ad Soyad", "Sınıf", "Doğum Tarihi"]


def test_hatalar_dosya_hazir_dedirtmez(ortam):
    t = fd.Tools()
    u = {"__user__": {"id": "u1", "role": "user"}}
    assert "bulunamadı" in asyncio.run(t.tablo_uret(__files__=[], **u))
    kayit, _, tmp = ortam
    yol = tmp / "a.csv"
    yol.write_text("Ad\nx\n", encoding="utf-8")
    kayit["f1"] = Dosya("a.csv", yol)
    f = {"__files__": [{"id": "f1", "type": "file"}], **u}
    assert "SÖYLEME" in asyncio.run(t.tablo_uret(cikar="yok", **f))
    assert "hiç satır kalmadı" in asyncio.run(t.tablo_uret(filtre_sutun="ad", filtre_deger="zzz", **f))
    assert "xlsx, docx ya da pdf" in asyncio.run(t.tablo_uret(bicim="exe", **f))


def test_pdf_girdisi_acik_reddedilir(ortam):
    kayit, _, tmp = ortam
    yol = tmp / "a.pdf"
    yol.write_bytes(b"%PDF")
    kayit["f1"] = Dosya("a.pdf", yol)
    s = asyncio.run(fd.Tools().tablo_bilgisi(__files__=[{"id": "f1"}], __user__={"id": "u1"}))
    assert "PDF'ten tablo okunamıyor" in s


def test_sohbeti_dosyaya_cevir_excel_her_tablo_ayri_sayfa(ortam):
    pytest.importorskip("openpyxl")
    _, kaydedilen, _ = ortam
    md = "| A | B |\n|---|---|\n| 1 | 2 |\n\nara\n\n| C |\n|---|\n| 3 |\n"
    c = asyncio.run(fd.Tools().sohbeti_dosyaya_cevir(md, "Rapor", "xlsx", __user__={"id": "u1"}))
    assert "Rapor.xlsx" in c
    from openpyxl import load_workbook

    assert load_workbook(io.BytesIO(kaydedilen[0][2])).sheetnames == ["Tablo 1", "Tablo 2"]
    assert "Metin boş" in asyncio.run(fd.Tools().sohbeti_dosyaya_cevir(" ", __user__={"id": "u1"}))
