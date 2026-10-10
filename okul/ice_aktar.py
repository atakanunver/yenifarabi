"""Excel aktarımı: önizleme (doğrulama) → taslak → onayla uygula. Hiçbir şeyi silmez."""

import io
import json
import re
import secrets
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime, timedelta

import auth
import openpyxl
import zaman
from metin import ilk_sifre, normalize_telefon, tr_kucuk

OGRENCI_BASLIK = [
    "okul_no",
    "ad_soyad",
    "sinif",
    "veli_ad",
    "veli_telefon",
    "veli2_ad",
    "veli2_telefon",
]
OGRETMEN_BASLIK = ["ad_soyad", "kullanici_adi", "telefon", "sinif", "ders"]
TASLAK_OMRU = timedelta(hours=1)
_SINIF_RE = re.compile(r"(9|10|11|12)-[A-ZÇĞİÖŞÜ]")


class AktarimHatasi(ValueError):
    pass


@dataclass
class Onizleme:
    satirlar: list[dict] = field(default_factory=list)
    hatalar: list[str] = field(default_factory=list)


@dataclass
class Sonuc:
    eklenen: int = 0
    guncellenen: int = 0
    yeni_hesaplar: list[tuple[str, str, str]] = field(
        default_factory=list
    )  # ad, kullanıcı adı, ilk şifre
    yeni_veliler: list[tuple[str, str]] = field(default_factory=list)  # ad, telefon
    excelde_olmayan: list[str] = field(default_factory=list)


def sablon(baslik: list[str]) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.append(baslik)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _hucre(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _oku(veri: bytes, baslik: list[str]) -> list[tuple[int, dict]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    except Exception as e:  # openpyxl bozuk dosyada çeşitli istisnalar atar
        raise AktarimHatasi(
            "Dosya okunamadı; .xlsx biçiminde bir Excel dosyası yükleyin."
        ) from e
    satirlar = list(wb.active.iter_rows(values_only=True))
    if not satirlar:
        raise AktarimHatasi("Dosya boş.")
    basliklar = [tr_kucuk(_hucre(h)) for h in satirlar[0]]
    eksik = [b for b in baslik if b not in basliklar]
    if eksik:
        raise AktarimHatasi(
            f"Eksik sütun(lar): {', '.join(eksik)}. Şablonu indirip kullanın."
        )
    sonuc = []
    for no, satir in enumerate(satirlar[1:], start=2):
        d = {
            b: _hucre(satir[basliklar.index(b)])
            if basliklar.index(b) < len(satir)
            else ""
            for b in baslik
        }
        if any(d.values()):
            sonuc.append((no, d))
    return sonuc


def _sinif(s: str) -> str | None:
    s = s.replace(" ", "").upper()
    return s if _SINIF_RE.fullmatch(s) else None


def ogrenci_onizle(veri: bytes) -> Onizleme:
    o = Onizleme()
    gorulen: set[int] = set()
    for no, d in _oku(veri, OGRENCI_BASLIK):
        hatalar = []
        try:
            okul_no = int(d["okul_no"])
        except ValueError:
            okul_no = None
            hatalar.append("okul_no sayı olmalı")
        if okul_no in gorulen:
            hatalar.append(f"okul_no {okul_no} dosyada tekrar ediyor")
        if not d["ad_soyad"]:
            hatalar.append("ad_soyad boş")
        sinif = _sinif(d["sinif"])
        if not sinif:
            hatalar.append(f"sınıf '{d['sinif']}' geçersiz (örnek: 9-A)")
        veliler = []
        for ad_k, tel_k in (("veli_ad", "veli_telefon"), ("veli2_ad", "veli2_telefon")):
            if d[tel_k]:
                tel = normalize_telefon(d[tel_k])
                if tel:
                    veliler.append([d[ad_k] or "Veli", tel])
                else:
                    hatalar.append(f"{tel_k} '{d[tel_k]}' geçersiz cep telefonu")
        if hatalar:
            o.hatalar.append(f"Satır {no}: " + "; ".join(hatalar))
            continue
        gorulen.add(okul_no)
        o.satirlar.append(
            {
                "okul_no": okul_no,
                "ad_soyad": d["ad_soyad"],
                "sinif": sinif,
                "veliler": veliler,
            }
        )
    return o


def _kullanici_ekle(conn, rol, ad, kullanici_adi=None, telefon=None, sifre=None) -> int:
    return conn.execute(
        "INSERT INTO kullanici (rol, ad_soyad, kullanici_adi, telefon, sifre_hash, sifre_degismeli, olusturma)"
        " VALUES (?, ?, ?, ?, ?, ?, ?)",
        (
            rol,
            ad,
            kullanici_adi,
            telefon,
            auth.sifre_hashle(sifre) if sifre else None,
            1 if sifre else 0,
            zaman.simdi_str(),
        ),
    ).lastrowid


def ogrenci_uygula(conn: sqlite3.Connection, satirlar: list[dict]) -> Sonuc:
    s = Sonuc()
    try:
        for r in satirlar:
            mevcut = conn.execute(
                "SELECT * FROM ogrenci WHERE okul_no = ?", (r["okul_no"],)
            ).fetchone()
            if mevcut:
                conn.execute(
                    "UPDATE ogrenci SET ad_soyad = ?, sinif = ? WHERE id = ?",
                    (r["ad_soyad"], r["sinif"], mevcut["id"]),
                )
                conn.execute(
                    "UPDATE kullanici SET ad_soyad = ? WHERE id = ?",
                    (r["ad_soyad"], mevcut["kullanici_id"]),
                )
                ogrenci_id = mevcut["id"]
                s.guncellenen += 1
            else:
                sifre = ilk_sifre(r["ad_soyad"])
                kid = _kullanici_ekle(
                    conn, "ogrenci", r["ad_soyad"], str(r["okul_no"]), sifre=sifre
                )
                ogrenci_id = conn.execute(
                    "INSERT INTO ogrenci (okul_no, ad_soyad, sinif, kullanici_id) VALUES (?, ?, ?, ?)",
                    (r["okul_no"], r["ad_soyad"], r["sinif"], kid),
                ).lastrowid
                s.eklenen += 1
                s.yeni_hesaplar.append((r["ad_soyad"], str(r["okul_no"]), sifre))
            for ad, tel in r["veliler"]:
                veli = conn.execute(
                    "SELECT id FROM kullanici WHERE rol = 'veli' AND telefon = ?",
                    (tel,),
                ).fetchone()
                if veli:
                    veli_id = veli["id"]
                else:
                    veli_id = _kullanici_ekle(conn, "veli", ad, telefon=tel)
                    s.yeni_veliler.append((ad, tel))
                conn.execute(
                    "INSERT OR IGNORE INTO veli_ogrenci (veli_id, ogrenci_id) VALUES (?, ?)",
                    (veli_id, ogrenci_id),
                )
        dosyadaki = {r["okul_no"] for r in satirlar}
        s.excelde_olmayan = [
            f"{o['okul_no']} {o['ad_soyad']} ({o['sinif']})"
            for o in conn.execute("SELECT * FROM ogrenci ORDER BY okul_no")
            if o["okul_no"] not in dosyadaki
        ]
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return s


def ogretmen_onizle(veri: bytes) -> Onizleme:
    o = Onizleme()
    for no, d in _oku(veri, OGRETMEN_BASLIK):
        hatalar = []
        if not d["ad_soyad"]:
            hatalar.append("ad_soyad boş")
        kadi = d["kullanici_adi"].lower()
        if not re.fullmatch(r"[a-z0-9._-]{3,30}", kadi):
            hatalar.append("kullanici_adi 3-30 karakter, yalnızca a-z 0-9 . _ - olmalı")
        tel = normalize_telefon(d["telefon"]) if d["telefon"] else None
        if d["telefon"] and not tel:
            hatalar.append(f"telefon '{d['telefon']}' geçersiz")
        sinif = _sinif(d["sinif"]) if d["sinif"] else None
        if d["sinif"] and not sinif:
            hatalar.append(f"sınıf '{d['sinif']}' geçersiz")
        if sinif and not d["ders"]:
            hatalar.append("sınıf verilmiş ama ders boş")
        if hatalar:
            o.hatalar.append(f"Satır {no}: " + "; ".join(hatalar))
            continue
        o.satirlar.append(
            {
                "ad_soyad": d["ad_soyad"],
                "kullanici_adi": kadi,
                "telefon": tel,
                "sinif": sinif,
                "ders": tr_kucuk(d["ders"]),
            }
        )
    return o


def ogretmen_uygula(conn: sqlite3.Connection, satirlar: list[dict]) -> Sonuc:
    s = Sonuc()
    try:
        for r in satirlar:
            mevcut = conn.execute(
                "SELECT * FROM kullanici WHERE kullanici_adi = ?", (r["kullanici_adi"],)
            ).fetchone()
            if mevcut and mevcut["rol"] not in ("ogretmen", "yonetici"):
                raise AktarimHatasi(
                    f"{r['kullanici_adi']} kullanıcı adı başka bir rolde kullanılıyor."
                )
            if mevcut:
                kid = mevcut["id"]
                if (mevcut["ad_soyad"], mevcut["telefon"]) != (
                    r["ad_soyad"],
                    r["telefon"],
                ):
                    conn.execute(
                        "UPDATE kullanici SET ad_soyad = ?, telefon = ? WHERE id = ?",
                        (r["ad_soyad"], r["telefon"], kid),
                    )
                    s.guncellenen += 1
            else:
                sifre = ilk_sifre(r["ad_soyad"])
                kid = _kullanici_ekle(
                    conn,
                    "ogretmen",
                    r["ad_soyad"],
                    r["kullanici_adi"],
                    r["telefon"],
                    sifre,
                )
                s.eklenen += 1
                s.yeni_hesaplar.append((r["ad_soyad"], r["kullanici_adi"], sifre))
            if r["sinif"]:
                conn.execute(
                    "INSERT OR IGNORE INTO ogretmen_gorev (ogretmen_id, sinif, ders) VALUES (?, ?, ?)",
                    (kid, r["sinif"], r["ders"]),
                )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return s


def taslak_kaydet(conn: sqlite3.Connection, tur: str, satirlar: list[dict]) -> str:
    tid = secrets.token_urlsafe(12)
    conn.execute(
        "INSERT INTO aktarim_taslak (id, tur, veri, olusturma) VALUES (?, ?, ?, ?)",
        (tid, tur, json.dumps(satirlar, ensure_ascii=False), zaman.simdi_str()),
    )
    conn.commit()
    return tid


def taslak_al(conn: sqlite3.Connection, tid: str, tur: str) -> list[dict] | None:
    r = conn.execute(
        "SELECT * FROM aktarim_taslak WHERE id = ? AND tur = ?", (tid, tur)
    ).fetchone()
    if (
        r is None
        or zaman.simdi() - datetime.strptime(r["olusturma"], zaman.BICIM) > TASLAK_OMRU
    ):
        return None
    return json.loads(r["veri"])


def taslak_sil(conn: sqlite3.Connection, tid: str) -> None:
    conn.execute("DELETE FROM aktarim_taslak WHERE id = ?", (tid,))
    conn.commit()
