"""Sınıf Excel ve PDF dosyalarından öğrenci doğum tarihleri ve veli/öğrenci
bağlantılı telefonlarını veritabanına aktarma ve karşılaştırmalı doğrulama scripti.

Kullanım:
    venv/bin/python scripts/sinif_bilgi_ice_aktar.py --kuru-calistir
    venv/bin/python scripts/sinif_bilgi_ice_aktar.py
"""

from __future__ import annotations

import argparse
from datetime import datetime
import io
import os
from pathlib import Path
import re
import shutil
import sqlite3
import sys
from zoneinfo import ZoneInfo

import openpyxl

# Proje kök dizinini sys.path'e ekle
KOK_DIZIN = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK_DIZIN))

import db
import gonderim

_ISTANBUL = ZoneInfo("Europe/Istanbul")
VARSAYILAN_SINIF_DIZINI = Path("/home/ata/farabi/mudur/SINIF")


def parse_dob_to_iso(val: str | None) -> str | None:
    """Tarih değerini ISO YYYY-MM-DD formatına çevirir."""
    if not val:
        return None
    s = str(val).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        return s
    m = re.match(r"^(\d{1,2})[/.](\d{1,2})[/.](\d{4})$", s)
    if m:
        d, mth, y = m.groups()
        return f"{int(y):04d}-{int(mth):02d}-{int(d):02d}"
    return None


def clean_phone(val: str | None) -> str | None:
    """Telefonu temizler, normalize eder ve doğrular."""
    if not val:
        return None
    s = str(val).strip()
    if not s or s.upper() == "YOK":
        return None
    p = gonderim.normalize_phone(s)
    return p if gonderim.is_valid_phone(p) else None


def dogum_tarihlerini_oku(sinif_dizini: Path) -> dict[str, dict]:
    """9-LAR.xlsx, 10-LAR.xlsx, 11ler.xlsx, 12ler.xlsx dosyalarından
    öğrenci doğum tarihlerini okur."""
    dosyalar = ["9-LAR.xlsx", "10-LAR.xlsx", "11ler.xlsx", "12ler.xlsx"]
    sonuclar = {}

    for dosya_adi in dosyalar:
        dosya_yolu = sinif_dizini / dosya_adi
        if not dosya_yolu.is_file():
            continue

        wb = openpyxl.load_workbook(dosya_yolu, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))

        for r in rows[1:]:
            if not r or not r[1]:
                continue
            ad_soyad = str(r[1]).strip()
            tc = str(r[2]).strip() if r[2] else None
            okul_no = r[3]
            sinif_str = str(r[4]).strip() if r[4] else None
            dob_iso = parse_dob_to_iso(r[5])
            norm = db._tr_norm(ad_soyad)

            sonuclar[norm] = {
                "ad_soyad": ad_soyad,
                "tc": tc,
                "okul_no": okul_no,
                "sinif_str": sinif_str,
                "dogum_tarihi": dob_iso,
                "kaynak_dosya": dosya_adi,
            }
    return sonuclar


def sinif_telefonlarini_oku(sinif_dizini: Path) -> list[dict]:
    """10-A.xlsx ve 12-A.xlsx dosyalarından öğrenci ve veli telefonlarını okur."""
    dosyalar = ["10-A.xlsx", "12-A.xlsx"]
    kayitlar = []

    for dosya_adi in dosyalar:
        dosya_yolu = sinif_dizini / dosya_adi
        if not dosya_yolu.is_file():
            continue

        wb = openpyxl.load_workbook(dosya_yolu, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))

        for r in rows[2:]:
            if not any(r):
                continue
            okul_no = r[0]
            veli_yakinlik = str(r[1]).strip() if r[1] else "Veli"
            veli_adi = str(r[2]).strip() if r[2] else ""
            ogr_ad = str(r[3]).strip() if r[3] else ""
            ogr_tel = clean_phone(r[4])
            anne_tel = clean_phone(r[5])
            baba_tel = clean_phone(r[6])

            if not ogr_ad:
                continue

            kayitlar.append({
                "kaynak_dosya": dosya_adi,
                "okul_no": okul_no,
                "veli_yakinlik": veli_yakinlik,
                "veli_adi": veli_adi,
                "ogr_ad": ogr_ad,
                "ogr_tel": ogr_tel,
                "anne_tel": anne_tel,
                "baba_tel": baba_tel,
            })
    return kayitlar


def pdf_dogrulama_yap(sinif_dizini: Path, excel_dobs: dict[str, dict]) -> list[dict]:
    """9A.PDF ve 9B.PDF dosyalarını okuyarak Excel doğum tarihleriyle
    karşılaştırmalı doğrulama yapar."""
    raporlar = []
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz
        except ImportError:
            return [{"mesaj": "pymupdf/fitz kütüphanesi bulunamadı, PDF doğrulama atlandı."}]

    for pdf_adi in ["9A.PDF", "9B.PDF"]:
        pdf_yolu = sinif_dizini / pdf_adi
        if not pdf_yolu.is_file():
            continue

        doc = fitz.open(pdf_yolu)
        for i, page in enumerate(doc):
            text = page.get_text()
            parts = re.split(r"I\.\s*Dönem\s*", text, maxsplit=1)
            name_parts = []
            okul_no = None
            if len(parts) > 1:
                lines = [l.strip() for l in parts[1].split("\n") if l.strip()]
                for l in lines:
                    if re.match(r"^\d+$", l):
                        okul_no = l
                        break
                    elif "Anne Adı" in l or "Baba Adı" in l:
                        break
                    else:
                        name_parts.append(l)
            ad = " ".join(name_parts).strip()

            m_dob = re.search(r"Doğum Tarihi\s*\n\s*:\s*(\d{2}/\d{2}/\d{4})", text)
            dob_raw = m_dob.group(1) if m_dob else None
            dob_iso = parse_dob_to_iso(dob_raw)

            # Yaş hesabı: doğum yılından 2026 yılına göre veya metindeki Yaşı alanı
            m_age = re.search(r"Yaşı\s*\n\s*:\s*(\d+)", text)
            if m_age:
                yas = int(m_age.group(1))
            elif dob_iso:
                yas = 2026 - int(dob_iso.split("-")[0])
            else:
                yas = None

            norm = db._tr_norm(ad)
            in_excel = excel_dobs.get(norm)

            durum = "eslesmedi"
            eslesme_detay = ""
            if in_excel:
                if in_excel["dogum_tarihi"] == dob_iso:
                    durum = "tam_eslesme"
                    eslesme_detay = f"9-LAR.xlsx ile birebir eşleşti ({dob_iso})"
                else:
                    durum = "tarih_uyusmazligi"
                    eslesme_detay = f"PDF: {dob_iso} != Excel: {in_excel['dogum_tarihi']}"
            else:
                durum = "listede_yok"
                eslesme_detay = f"Yaş {yas} (Aktif öğrenci listesinde bulunmuyor, hariç tutuldu)"

            raporlar.append({
                "pdf": pdf_adi,
                "sayfa": i + 1,
                "ad_soyad": ad,
                "okul_no": okul_no,
                "dob_iso": dob_iso,
                "yas": yas,
                "durum": durum,
                "detay": eslesme_detay,
            })
    return raporlar


def analiz_et(sinif_dizini: Path = VARSAYILAN_SINIF_DIZINI, conn: sqlite3.Connection | None = None) -> dict:
    """Tüm sınıf dosyalarını okuyarak veritabanı ile karşılaştırır ve
    yapılacak işlemleri çıkarır."""
    kendi_baglanti = False
    if conn is None:
        conn = db.baglanti()
        kendi_baglanti = True

    try:
        # Aktif öğrencileri çek
        ogrenciler = conn.execute(
            "SELECT k.id, k.ad_soyad, s.ad AS sinif_ad, k.sinif_id, k.dogum_tarihi, k.telefon "
            "FROM kisiler k JOIN siniflar s ON s.id = k.sinif_id "
            "WHERE k.tur = 'ogrenci'"
        ).fetchall()

        db_ogrenciler = {db._tr_norm(ogr["ad_soyad"]): dict(ogr) for ogr in ogrenciler}

        # Mevcut velileri çek
        veliler = conn.execute(
            "SELECT k.id, k.ad_soyad, k.telefon, k.sinif_id, k.ogrenci_kisi_id "
            "FROM kisiler k WHERE k.tur = 'veli'"
        ).fetchall()
        db_veliler = [dict(v) for v in veliler]

        excel_dobs = dogum_tarihlerini_oku(sinif_dizini)
        telefon_kayitlari = sinif_telefonlarini_oku(sinif_dizini)
        pdf_raporu = pdf_dogrulama_yap(sinif_dizini, excel_dobs)

        # 1. Doğum Tarihi Analizi
        dogum_guncellemeleri = []
        dogum_dogrulananlar = []
        dogum_eslesmeyen_db = []

        for norm, ogr in db_ogrenciler.items():
            if norm in excel_dobs:
                ex_info = excel_dobs[norm]
                yeni_dob = ex_info["dogum_tarihi"]
                mevcut_dob = ogr["dogum_tarihi"]

                if not mevcut_dob and yeni_dob:
                    dogum_guncellemeleri.append({
                        "id": ogr["id"],
                        "ad_soyad": ogr["ad_soyad"],
                        "sinif_ad": ogr["sinif_ad"],
                        "eski_dob": None,
                        "yeni_dob": yeni_dob,
                        "kaynak": ex_info["kaynak_dosya"],
                    })
                elif mevcut_dob == yeni_dob:
                    dogum_dogrulananlar.append({
                        "id": ogr["id"],
                        "ad_soyad": ogr["ad_soyad"],
                        "sinif_ad": ogr["sinif_ad"],
                        "dob": mevcut_dob,
                    })
                else:
                    dogum_guncellemeleri.append({
                        "id": ogr["id"],
                        "ad_soyad": ogr["ad_soyad"],
                        "sinif_ad": ogr["sinif_ad"],
                        "eski_dob": mevcut_dob,
                        "yeni_dob": yeni_dob,
                        "kaynak": ex_info["kaynak_dosya"],
                    })
            else:
                dogum_eslesmeyen_db.append({
                    "id": ogr["id"],
                    "ad_soyad": ogr["ad_soyad"],
                    "sinif_ad": ogr["sinif_ad"],
                    "mevcut_dob": ogr["dogum_tarihi"],
                })

        # 2. Öğrenci Telefon Güncellemeleri
        ogrenci_tel_guncellemeleri = []
        veli_islemleri = []

        for tk in telefon_kayitlari:
            norm = db._tr_norm(tk["ogr_ad"])
            ogr = db_ogrenciler.get(norm)
            if not ogr:
                continue

            ogr_id = ogr["id"]
            sinif_id = ogr["sinif_id"]

            # Öğrenci telefonu
            if tk["ogr_tel"]:
                if ogr["telefon"] != tk["ogr_tel"]:
                    ogrenci_tel_guncellemeleri.append({
                        "id": ogr_id,
                        "ad_soyad": ogr["ad_soyad"],
                        "sinif_ad": ogr["sinif_ad"],
                        "eski_tel": ogr["telefon"],
                        "yeni_tel": tk["ogr_tel"],
                    })

            # Veliler
            yakinlik = tk["veli_yakinlik"].upper()
            v_adi = tk["veli_adi"]
            anne_tel = tk["anne_tel"]
            baba_tel = tk["baba_tel"]

            # Birincil ve ikincil veli belirleme
            if "ANNE" in yakinlik:
                birincil_ad = v_adi or f"{ogr['ad_soyad']} Annesi"
                birincil_tel = anne_tel
                ikincil_ad = f"{ogr['ad_soyad']} Babası"
                ikincil_tel = baba_tel
            elif "BABA" in yakinlik:
                birincil_ad = v_adi or f"{ogr['ad_soyad']} Babası"
                birincil_tel = baba_tel
                ikincil_ad = f"{ogr['ad_soyad']} Annesi"
                ikincil_tel = anne_tel
            else:
                birincil_ad = v_adi or f"{ogr['ad_soyad']} Velisi"
                birincil_tel = anne_tel or baba_tel
                ikincil_ad = f"{ogr['ad_soyad']} ({tk['veli_yakinlik']})"
                ikincil_tel = baba_tel if birincil_tel == anne_tel else anne_tel

            adaylar = []
            if birincil_tel:
                adaylar.append((birincil_ad, birincil_tel, "Birincil Veli"))
            if ikincil_tel and ikincil_tel != birincil_tel:
                adaylar.append((ikincil_ad, ikincil_tel, "İkincil Veli / İletişim"))

            for ad, tel, rol in adaylar:
                # Mevcut veliler arasında var mı kontrol et (öğrenciye bağlı + telefon veya isim)
                var_olan = None
                for mv in db_veliler:
                    if mv["ogrenci_kisi_id"] == ogr_id and (
                        mv["telefon"] == tel or db._tr_norm(mv["ad_soyad"]) == db._tr_norm(ad)
                    ):
                        var_olan = mv
                        break

                if var_olan:
                    if var_olan["telefon"] != tel or var_olan["ad_soyad"] != ad:
                        veli_islemleri.append({
                            "islem": "guncelle",
                            "kisi_id": var_olan["id"],
                            "ad_soyad": ad,
                            "telefon": tel,
                            "sinif_id": sinif_id,
                            "ogrenci_kisi_id": ogr_id,
                            "ogrenci_ad": ogr["ad_soyad"],
                            "rol": rol,
                        })
                else:
                    veli_islemleri.append({
                        "islem": "ekle",
                        "kisi_id": None,
                        "ad_soyad": ad,
                        "telefon": tel,
                        "sinif_id": sinif_id,
                        "ogrenci_kisi_id": ogr_id,
                        "ogrenci_ad": ogr["ad_soyad"],
                        "rol": rol,
                    })

        return {
            "toplam_db_ogrenci": len(db_ogrenciler),
            "dogum_dogrulanan_adet": len(dogum_dogrulananlar),
            "dogum_guncellenecek": dogum_guncellemeleri,
            "dogum_eslesmeyen_db": dogum_eslesmeyen_db,
            "ogrenci_tel_guncellenecek": ogrenci_tel_guncellemeleri,
            "veli_islemleri": veli_islemleri,
            "pdf_raporu": pdf_raporu,
        }
    finally:
        if kendi_baglanti:
            conn.close()


def yedek_olustur() -> Path:
    """Veritabanının tarihli yedeğini oluşturur."""
    db_yolu = db.DB_YOLU
    yedek_dizin = db_yolu.parent / "yedek"
    yedek_dizin.mkdir(parents=True, exist_ok=True)
    zaman_damgasi = datetime.now(_ISTANBUL).strftime("%Y-%m-%d_%H%M%S")
    yedek_dosya = yedek_dizin / f"smssistemi_{zaman_damgasi}_sinif_bilgi_oncesi.db"
    shutil.copy2(db_yolu, yedek_dosya)
    return yedek_dosya


def aktarim_uygula(analiz_sonucu: dict, conn: sqlite3.Connection) -> dict:
    """Analiz sonucundaki işlemleri atomik transaction içinde veritabanına uygular."""
    cur = conn.cursor()

    dogum_guncellendi = 0
    ogrenci_tel_guncellendi = 0
    veli_eklendi = 0
    veli_guncellendi = 0

    # 1. Öğrenci Doğum Tarihleri
    for d in analiz_sonucu["dogum_guncellenecek"]:
        cur.execute(
            "UPDATE kisiler SET dogum_tarihi = ? WHERE id = ? AND tur = 'ogrenci'",
            (d["yeni_dob"], d["id"]),
        )
        dogum_guncellendi += 1

    # 2. Öğrenci Telefonları
    for t in analiz_sonucu["ogrenci_tel_guncellenecek"]:
        cur.execute(
            "UPDATE kisiler SET telefon = ? WHERE id = ? AND tur = 'ogrenci'",
            (t["yeni_tel"], t["id"]),
        )
        ogrenci_tel_guncellendi += 1

    # 3. Veliler (Deduplicated insert/update)
    for v in analiz_sonucu["veli_islemleri"]:
        if v["islem"] == "ekle":
            cur.execute(
                "INSERT INTO kisiler (ad_soyad, telefon, sinif_id, tur, ogrenci_kisi_id) "
                "VALUES (?, ?, ?, 'veli', ?)",
                (v["ad_soyad"], v["telefon"], v["sinif_id"], v["ogrenci_kisi_id"]),
            )
            veli_eklendi += 1
        elif v["islem"] == "guncelle":
            cur.execute(
                "UPDATE kisiler SET ad_soyad = ?, telefon = ?, sinif_id = ? "
                "WHERE id = ? AND tur = 'veli'",
                (v["ad_soyad"], v["telefon"], v["sinif_id"], v["kisi_id"]),
            )
            veli_guncellendi += 1

    conn.commit()

    return {
        "dogum_guncellendi": dogum_guncellendi,
        "ogrenci_tel_guncellendi": ogrenci_tel_guncellendi,
        "veli_eklendi": veli_eklendi,
        "veli_guncellendi": veli_guncellendi,
    }


def main() -> None:
    ayristirici = argparse.ArgumentParser(description=__doc__)
    ayristirici.add_argument(
        "--dizin",
        type=Path,
        default=VARSAYILAN_SINIF_DIZINI,
        help=f"Sınıf dosyaları dizini (varsayılan: {VARSAYILAN_SINIF_DIZINI})",
    )
    ayristirici.add_argument(
        "--kuru-calistir",
        action="store_true",
        help="DB'ye yazmaz, sadece karşılaştırmalı doğrulama analizini basar.",
    )
    ayristirici.add_argument(
        "--yedek-alma",
        action="store_true",
        help="İşlem öncesi veritabanı yedeği almayı atlar.",
    )
    args = ayristirici.parse_args()

    if not args.dizin.is_dir():
        print(f"Hata: Dizin bulunamadı: {args.dizin}", file=sys.stderr)
        sys.exit(1)

    db.semayi_kur()
    conn = db.baglanti()
    try:
        analiz = analiz_et(args.dizin, conn=conn)

        print("=" * 70)
        print("  FARABİ ANADOLU LİSESİ - SINIF BİLGİLERİ DOĞRULAMA VE AKTARIM")
        print("=" * 70)
        print(f"Kaynak Dizin: {args.dizin}")
        print(f"Veritabanındaki Aktif Öğrenci Sayısı: {analiz['toplam_db_ogrenci']}")
        print(f"• Mevcut ve Birebir Doğrulanan Doğum Tarihleri: {analiz['dogum_dogrulanan_adet']}")
        print(f"• Yeni Eklenecek/Güncellenecek Doğum Tarihleri: {len(analiz['dogum_guncellenecek'])}")
        print(f"• Dosyalarda Bulunmayan Öğrenciler: {len(analiz['dogum_eslesmeyen_db'])}")
        for e in analiz["dogum_eslesmeyen_db"]:
            print(f"    - [{e['sinif_ad']}] {e['ad_soyad']} (Mevcut Doğum Tarihi: {e['mevcut_dob'] or 'YOK'})")
        print(f"• Güncellenecek Öğrenci Telefonları: {len(analiz['ogrenci_tel_guncellenecek'])}")
        print(f"• İşlenecek Veli Kayıtları: {len(analiz['veli_islemleri'])}")

        eklenecek_veli = sum(1 for v in analiz["veli_islemleri"] if v["islem"] == "ekle")
        guncellenecek_veli = sum(1 for v in analiz["veli_islemleri"] if v["islem"] == "guncelle")
        print(f"    - Yeni Eklenecek: {eklenecek_veli}")
        print(f"    - Var Olan Güncellenecek: {guncellenecek_veli}")

        # PDF Karşılaştırma Özeti
        pdf_eslesen = sum(1 for p in analiz["pdf_raporu"] if p.get("durum") == "tam_eslesme")
        pdf_eski = sum(1 for p in analiz["pdf_raporu"] if p.get("durum") == "listede_yok")
        print(f"\nPDF Doğrulama Özeti (9A.PDF & 9B.PDF):")
        print(f"  • Aktif Öğrencilerle Birebir Eşleşen: {pdf_eslesen}")
        print(f"  • Kurumda Olmayan / Eski Kayıt (Atlanan): {pdf_eski}")

        if args.kuru_calistir:
            print("\n" + "-" * 70)
            print("[KURU ÇALIŞTIRMA — DB'ye hiçbir kayıt yazılmadı]")
            print("-" * 70)
            print("\nÖrnek Doğum Tarihi Güncellemeleri:")
            for d in analiz["dogum_guncellenecek"][:5]:
                print(f"  [{d['sinif_ad']}] {d['ad_soyad']}: {d['eski_dob']} -> {d['yeni_dob']} ({d['kaynak']})")
            print(f"  ... ve {max(0, len(analiz['dogum_guncellenecek']) - 5)} öğrenci daha.")

            print("\nÖrnek Öğrenci Telefon Güncellemeleri:")
            for t in analiz["ogrenci_tel_guncellenecek"][:5]:
                print(f"  [{t['sinif_ad']}] {t['ad_soyad']}: {t['eski_tel']} -> {t['yeni_tel']}")

            print("\nÖrnek Veli Kayıtları:")
            for v in analiz["veli_islemleri"][:8]:
                print(f"  [{v['islem'].upper()}] {v['ad_soyad']} ({v['rol']}) - Tel: {v['telefon']} -> Öğrenci: {v['ogrenci_ad']}")
            return

        # Gerçek Aktarım
        if not args.yedek_alma:
            yedek_yolu = yedek_olustur()
            print(f"\nDB Yedeği Alındı: {yedek_yolu.name}")

        sonuc = aktarim_uygula(analiz, conn)

        print("\n" + "=" * 70)
        print("  AKTARIM BAŞARIYLA TAMAMLANDI")
        print("=" * 70)
        print(f"• Güncellenen Öğrenci Doğum Tarihleri: {sonuc['dogum_guncellendi']}")
        print(f"• Güncellenen Öğrenci Telefonları: {sonuc['ogrenci_tel_guncellendi']}")
        print(f"• Eklenen Yeni Veli Kayıtları: {sonuc['veli_eklendi']}")
        print(f"• Güncellenen Veli Kayıtları: {sonuc['veli_guncellendi']}")

        # Son Durum Doğrulaması
        toplam_kisi = conn.execute("SELECT count(*) FROM kisiler").fetchone()[0]
        tur_sayilari = dict(conn.execute("SELECT tur, count(*) FROM kisiler GROUP BY tur").fetchall())
        dogum_sayilari = dict(conn.execute("SELECT tur, count(dogum_tarihi) FROM kisiler GROUP BY tur").fetchall())
        tel_sayilari = dict(conn.execute("SELECT tur, count(telefon) FROM kisiler GROUP BY tur").fetchall())

        print("\nVeritabanı Son Durum:")
        print(f"  Toplam Kişi: {toplam_kisi}")
        for t in ["ogrenci", "personel", "veli"]:
            sayi = tur_sayilari.get(t, 0)
            dt_sayi = dogum_sayilari.get(t, 0)
            tel_sayi = tel_sayilari.get(t, 0)
            print(f"  - {t.capitalize():8s}: {sayi:3d} kişi | Doğum Tarihi: {dt_sayi:3d} | Telefon: {tel_sayi:3d}")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
