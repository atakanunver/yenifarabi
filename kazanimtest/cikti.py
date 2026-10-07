"""kazanimtest/cikti.py — Öğretmen için Excel + Word (öğrenci kopyası + cevap anahtarı).
Yazım yeri repo değil: /mnt/farabi-data/farabi/kazanim_testleri/."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

HARF = "ABCD"


def dosya_tabani(dizin: str | Path, tarih, sinif: str, ders: str) -> Path:
    d = Path(dizin)
    d.mkdir(parents=True, exist_ok=True)
    return d / f"{tarih.isoformat()}_{sinif}_{ders}"


def excel_yaz(yol: Path, baslik: str, sorular: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test"
    ws.append([baslik])
    ws.append(["No", "Soru", "A", "B", "C", "D", "Doğru", "Kaynak"])
    for i, s in enumerate(sorular, 1):
        ws.append([i, s["soru"], *s["secenekler"], HARF[s["dogru_index"]], s.get("etiket", "")])
    ws.column_dimensions["B"].width = 80
    hedef = yol.with_suffix(".xlsx")
    wb.save(hedef)
    return hedef


def word_yaz(yol: Path, baslik: str, sorular: list[dict]) -> Path:
    d = Document()
    d.add_heading(baslik, level=1)
    d.add_paragraph("Ad Soyad: ____________________    Okul No: ________")
    for i, s in enumerate(sorular, 1):
        d.add_paragraph(f"{i}. {s['soru']}")
        for j, sec in enumerate(s["secenekler"]):
            d.add_paragraph(f"{HARF[j]}) {sec}")
    d.add_page_break()
    d.add_heading("Cevap Anahtarı", level=1)
    for i, s in enumerate(sorular, 1):
        d.add_paragraph(f"{i}. {HARF[s['dogru_index']]}")
    hedef = yol.with_suffix(".docx")
    d.save(hedef)
    return hedef
