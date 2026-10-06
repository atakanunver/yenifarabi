"""idari_yukle.py — saf fonksiyonlar. Gerçek tesseract/bge-m3/DB'ye gitmez."""

from pathlib import Path

import docx
import fitz
import idari_yukle as iy


def test_bilinen_dosya_adi_esleme():
    assert iy.belge_adi(Path("yazıı ve uygulamalı sınavlar yönergesi.pdf")) == \
        "Yazılı ve Uygulamalı Sınavlar Yönergesi"


def test_bilinmeyen_dosya_adi_dosya_kokunden():
    assert iy.belge_adi(Path("okul_nobet-cizelgesi.pdf")) == "Okul Nobet Cizelgesi"


def test_listeleme_alt_klasor_haric_xlsx_dahil(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"x")
    (tmp_path / "b.docx").write_bytes(b"x")
    (tmp_path / "c.jpeg").write_bytes(b"x")
    (tmp_path / "d.xlsx").write_bytes(b"x")
    (tmp_path / "GUVENLIK").mkdir()
    (tmp_path / "GUVENLIK" / "e.pdf").write_bytes(b"x")
    assert [p.name for p in iy.belgeleri_listele(tmp_path)] == ["a.pdf", "b.docx", "c.jpeg", "d.xlsx"]


def test_xlsx_sayfa_basina_satirlar(tmp_path):
    import openpyxl
    yol = tmp_path / "plan.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "EKİM"
    wb.active.append(["Hafta", "Etkinlik", None])
    wb.active.append([1, "Veli toplantısı", ""])
    wb.create_sheet("BOŞ")
    wb.save(yol)
    s = iy.sayfa_metinleri(yol)
    assert s[0] == (1, "EKİM\nHafta | Etkinlik\n1 | Veli toplantısı", False)
    assert s[1] == (2, "", False)


def test_haric_dosya_listelenmez(tmp_path):
    (tmp_path / "657DEVLET.pdf").write_bytes(b"x")
    (tmp_path / "657_devlet_memurlari_kanunu.pdf").write_bytes(b"x")
    assert [p.name for p in iy.belgeleri_listele(tmp_path)] == ["657DEVLET.pdf"]


def test_metinli_pdf_ocr_suz(tmp_path):
    yol = tmp_path / "m.pdf"
    d = fitz.open()
    d.new_page().insert_text((72, 72), "Madde 1 - Bu yonetmelik ortaogretim kurumlarini kapsar. " * 3)
    d.save(yol)
    cagri = []
    sonuc = iy.sayfa_metinleri(yol, ocr=lambda png: cagri.append(1) or "OCR")
    assert sonuc[0][0] == 1 and "Madde 1" in sonuc[0][1] and sonuc[0][2] is False
    assert cagri == []


def test_bos_pdf_sayfasi_ocr_a_gider(tmp_path):
    yol = tmp_path / "t.pdf"
    d = fitz.open()
    d.new_page()
    d.save(yol)
    sonuc = iy.sayfa_metinleri(yol, ocr=lambda png: "taranmış metin")
    assert sonuc == [(1, "taranmış metin", True)]


def test_docx_paragraf_ve_tablo(tmp_path):
    yol = tmp_path / "y.docx"
    belge = docx.Document()
    belge.add_paragraph("Sınıf rehber öğretmeni görevleri")
    t = belge.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Görev"
    t.cell(0, 1).text = "Veli toplantısı"
    belge.save(yol)
    (sayfa, metin, ocr_mu), = iy.sayfa_metinleri(yol)
    assert sayfa == 1 and "rehber" in metin and "Veli toplantısı" in metin and ocr_mu is False


def test_gorsel_dosya_ocr(tmp_path):
    yol = tmp_path / "g.jpeg"
    yol.write_bytes(b"\xff\xd8sahte")
    assert iy.sayfa_metinleri(yol, ocr=lambda b: "norm saat") == [(1, "norm saat", True)]


def test_kisisel_veri_dosyalari_haric_tutulur(tmp_path):
    (tmp_path / "bilgi fomu detay20262027.docx").write_bytes(b"x")
    (tmp_path / "personel bilgi2026.xlsx").write_bytes(b"x")
    (tmp_path / "mevzuat.pdf").write_bytes(b"x")
    listelenen = [p.name for p in iy.belgeleri_listele(tmp_path, haric=iy.KISISEL_VERI_DOSYALARI)]
    assert "bilgi fomu detay20262027.docx" not in listelenen
    assert "personel bilgi2026.xlsx" not in listelenen
    assert listelenen == ["mevzuat.pdf"]


def test_temizle_metin_soft_hyphen_ve_tire():
    ham = "fo\xad to\xad sen\xad tez ve mad-\n delerin tepkimesi \x00\x08\ufffd"
    temiz = iy.temizle_metin(ham)
    assert "\xad" not in temiz
    assert "\x00" not in temiz
    assert "\x08" not in temiz
    assert "\ufffd" not in temiz
    assert "maddelerin" in temiz


def test_belge_parcala_madde_baslikli():
    metin = (
        "Mazeret izni: Madde 104 – A) Kadın memura doğum izni verilir.\n"
        "B) Memura eşinin doğum yapması halinde babalık izni verilir.\n"
        "Hastalık izni: Madde 105 – Memura hastalık raporuna göre izin verilir."
    )
    sayfalar = [(1, metin, False)]
    chunks = iy.belge_parcala(sayfalar, "657 Sayılı Kanun")
    assert len(chunks) >= 2
    # İlk parça Madde 104 başlığını taşımalı
    assert "[657 Sayılı Kanun - Mazeret izni: Madde 104" in chunks[0][1]
    assert "Kadın memura" in chunks[0][1]
    # Madde 105 başlığı da ayrı parçada bulunmalı
    madde_105_chunk = next(c[1] for c in chunks if "Madde 105" in c[1])
    assert "[657 Sayılı Kanun - Hastalık izni: Madde 105" in madde_105_chunk


def test_parcala_kati_limit_asmaz():
    uzun = "Kelime " * 400
    parcalar = iy._parcala_kati(uzun, limit=200)
    assert len(parcalar) > 1
    assert all(len(p) <= 200 for p in parcalar)

